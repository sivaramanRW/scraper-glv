"""Scrape Rural District > Taluk > Village > Survey Number > Sub Division from Tamil Nilam GI Viewer, through its APIs.

Usage:
    pip install -r requirements.txt          # requests, cryptography, PySocks, openpyxl: no browser
    python scrape.py                         # everything
    python scrape.py --district Ariyalur     # one district
    python scrape.py --district Ariyalur --taluk Andimadam --village Aiyyur
    python scrape.py --worker 1/3            # one of 3 parallel workers (see below)

API client: tngis_api.py (encrypted envelope, anonymous session key, captcha + login, token refresh); the calls are listed in
API_REFERENCE.md. The old Chromium version is kept in backup/scrape_chromium.py.

Parallel workers: --worker k/N gives this process the accounts k, k+N, k+2N, ... and, inside every taluk, the villages whose
position in the listing is k, k+N, ... so N workers never share an account or a workbook. With 69 accounts and 3 workers each
worker has 23 accounts.

Accounts: a worker takes its accounts in turn. A *visit* = get a session for the account (saved tokens are refreshed, so a login
with captcha is only needed when the refresh token is dead; the captcha is read by the Qwen3-VL API fleet CAPTCHA_API_URLS), then
at most ACCOUNT_BUDGET (40) check-areg calls, --delay (5 s) apart, then on to the next account. The site's limits per account
are 50 calls/1 h, 125/4 h, 175/8 h and 250/24 h and it only reports the one that trips, so the script keeps the time of every
call (out/.progress/calls[.k]) and never exceeds CALL_CAPS (a margin below each limit: 62 min:40, 4 h 2 m:112, 8 h 2 m:157,
24 h 2 m:225). An account with less than ACCOUNT_MIN_VISIT calls of room is left alone; when no account has room the worker
sleeps until the first one has. A refusal the server does give ("Retry after ...") rests that account for the time it names
(out/.progress/cooldowns[.k]). The last account used is out/.progress/account_idx[.k].

Login lockout ("Too many failed login attempts from this network. Try again in 11hr 52min"): retrying during it restarts the
timer, so the worker sleeps the whole time without trying (the end time is in out/.progress/lockout, shared by all workers
behind the same IP, per worker with IPS_PROXY_METHOD=ONLINE_PROXY).

Proxy: IPS_PROXY_METHOD=ROTATION_IP (default) sends everything through PROXY_URL (e.g. socks5://192.168.59.64:1080, the phone
whose IP scripts/rotate_phone_ip.py changes); ONLINE_PROXY gives worker k the k-th of ONLINE_PROXY_URLS=url0,url1,url2.
The captcha API (LAN) is always called directly.

Output: out/Village_map_of_<District>_<Taluk>_<Village>_<lat>_<long>.xlsx   (columns: surveynumber | sub-divisions)
        lat/lon = centre of the village boundary's bounding box (one get-geom call per village).
        "-" in the sub-divisions column means the server reported no sub-divisions for that survey number.
Resumable / crash-safe, Excel only: the workbook is created with every survey number as soon as the village starts and is
        re-saved (atomically) after every fetched row. A blank sub-divisions cell = not fetched yet, so on restart the script
        opens the workbook and continues with the blank rows. out/ is committed to git: scripts/sync_progress.sh (pm2
        land-progress-sync, every 10 min) pushes it, so a fresh clone resumes where the last machine stopped. Run one machine at a time.
        Older files without lat/lon in the name are backfilled: the centre is fetched once and the file is renamed.
Captcha: POSTs {"images_base64": [<png>], "prompt": CAPTCHA_PROMPT, "max_new_tokens": 16, "temperature": 0} to
<url>/v1/describe/base64 and takes the "text" field (bearer token CAPTCHA_API_TOKEN, or VLM_API_TOKENS in CAPTCHA_API_TOKEN_FILE).
Worker k starts at API k and consecutive attempts rotate; an answer whose length != CAPTCHA_LEN (default 6) is discarded and a
fresh captcha is requested (nothing is submitted, so no failed login is counted). Without CAPTCHA_API_URLS a local Ollama model is used.
"""
import argparse, base64, json, os, re, sys, time, traceback, urllib.request
from pathlib import Path
import requests
from openpyxl import Workbook, load_workbook
from tngis_api import Api, AuthError, Locked, parse_wait

HERE = Path(__file__).parent
for _line in (HERE / ".env").read_text().splitlines() if (HERE / ".env").exists() else []:  # .env = defaults only
    if _line.strip() and not _line.lstrip().startswith("#") and "=" in _line:
        _k, _v = _line.split("=", 1); os.environ.setdefault(_k.strip(), _v.strip().strip("'\""))

ap = argparse.ArgumentParser()
ap.add_argument("--district"); ap.add_argument("--taluk"); ap.add_argument("--village")
ap.add_argument("--out", default=str(HERE / "out"))
ap.add_argument("--delay", type=float, default=float(os.environ.get("CALL_DELAY_S", "5")), help="seconds between check-areg calls")
ap.add_argument("--worker", default="0/1", help="k/N: this is worker k of N parallel workers (accounts and villages are split)")
args = ap.parse_args()
OUT = Path(args.out); (OUT / ".progress").mkdir(parents=True, exist_ok=True)
TOKEN_DIR = HERE / "state" / "tokens"  # access/refresh tokens per account: secrets, never inside out/ (which is committed)


def safe(s): return re.sub(r"[^\w.-]+", "_", s.strip()).strip("_")
def match(name, q): return not q or q.lower() in name.lower()


ACCOUNTS = json.load(open(HERE / "accounts.json"))
WK, NWK = (int(x) for x in args.worker.split("/"))
assert 0 <= WK < NWK, "--worker must be k/N with 0 <= k < N"
ACCOUNT_POOL = list(range(len(ACCOUNTS)))[WK::NWK]
if WK: ACCOUNT_POOL.reverse()  # extra workers start from the far end of their share, away from where a lone run left off
assert ACCOUNT_POOL, "more workers than accounts"
PROG = OUT / ".progress"
def pfile(name): return PROG / (name if WK == 0 else f"{name}.{WK}")


IDX_FILE, COOL_FILE, CALLS_FILE = pfile("account_idx"), pfile("cooldowns"), pfile("calls")
ACTIVE_WINDOW = 180  # s: a workbook saved more recently than this is being written by another process

# ---- proxy
METHOD = os.environ.get("IPS_PROXY_METHOD", "ROTATION_IP").strip().upper()
assert METHOD in ("ROTATION_IP", "ONLINE_PROXY"), "IPS_PROXY_METHOD must be ROTATION_IP or ONLINE_PROXY"
LOCK_FILE = PROG / ("lockout" if METHOD == "ROTATION_IP" else f"lockout.{WK}")  # ROTATION_IP: every worker shares one IP, so one lockout


def pick_proxy():
    if METHOD == "ONLINE_PROXY":
        urls = [u.strip() for u in os.environ.get("ONLINE_PROXY_URLS", "").split(",") if u.strip()]
        assert urls, "IPS_PROXY_METHOD=ONLINE_PROXY needs ONLINE_PROXY_URLS=url0,url1,url2 (worker k uses the k-th)"
        p = urls[WK % len(urls)]
    else:
        p = os.environ.get("PROXY_URL", "").strip()
    if not p: return None
    if "://" not in p: p = "http://" + p
    return re.sub(r"^socks5://", "socks5h://", p)  # socks5h: the proxy resolves DNS, like a browser would


PROXY = pick_proxy()

# ---- captcha reader
OLLAMA = os.environ.get("OLLAMA_URL", "http://localhost:11434")
VLM_MODELS = [m.strip() for m in os.environ.get("CAPTCHA_MODEL", "qwen2.5vl:7b").split(",") if m.strip()]
CAPTCHA_APIS = [u.strip().rstrip("/") for u in os.environ.get("CAPTCHA_API_URLS", "").split(",") if u.strip()]


def _captcha_token():
    """CAPTCHA_API_TOKEN, else the first token of VLM_API_TOKENS=... in CAPTCHA_API_TOKEN_FILE (the API's own .env)."""
    if os.environ.get("CAPTCHA_API_TOKEN"): return os.environ["CAPTCHA_API_TOKEN"]
    f = os.environ.get("CAPTCHA_API_TOKEN_FILE")
    if f and os.path.exists(f):
        for line in Path(f).read_text().splitlines():
            if line.startswith("VLM_API_TOKENS="): return line.split("=", 1)[1].split(",")[0].strip()
        return Path(f).read_text().strip().split(",")[0]  # plain token file
    return ""


CAPTCHA_API_TOKEN = _captcha_token()
CAPTCHA_LEN = int(os.environ.get("CAPTCHA_LEN", "6"))  # expected captcha length; 0 disables the check
CAPTCHA_PROMPT = os.environ.get("CAPTCHA_PROMPT", "This is a CAPTCHA image containing exactly 6 characters (letters and digits, "
                                "case-sensitive) with distracting lines. Reply with ONLY the 6 characters, nothing else.")
CAPTCHA_API_TIMEOUT = int(os.environ.get("CAPTCHA_API_TIMEOUT", "60"))  # s


def read_captcha_api(png, attempt=0):
    """Ask the Qwen3-VL API fleet for the captcha text. Returns '' if every API failed or the answer looks wrong."""
    body = json.dumps({"images_base64": [base64.b64encode(png).decode()], "prompt": CAPTCHA_PROMPT, "max_new_tokens": 16, "temperature": 0}).encode()
    hdr = {"Content-Type": "application/json", "Authorization": f"Bearer {CAPTCHA_API_TOKEN}"}
    start = (WK + attempt) % len(CAPTCHA_APIS)  # worker k prefers API k; later attempts rotate
    for i in range(len(CAPTCHA_APIS)):
        base = CAPTCHA_APIS[(start + i) % len(CAPTCHA_APIS)]
        url = base if base.endswith("/describe/base64") else base + "/v1/describe/base64"
        try:
            with urllib.request.urlopen(urllib.request.Request(url, body, hdr), timeout=CAPTCHA_API_TIMEOUT) as r:
                d = json.load(r)
            txt = re.sub(r"\W+", "", str(d.get("text", "")))
            if CAPTCHA_LEN and len(txt) != CAPTCHA_LEN:
                print(f"  captcha API {base}: answer {txt!r} is not {CAPTCHA_LEN} chars; asking for a fresh captcha"); return ""
            return txt
        except Exception as e:
            print(f"  captcha API {base} error: {str(e).splitlines()[0][:120]}")
    return ""


def read_captcha(png, attempt=0):
    """Captcha text via the API fleet (CAPTCHA_API_URLS) or, if none is configured, a local Ollama model. '' = unreadable."""
    if CAPTCHA_APIS: return read_captcha_api(png, attempt)
    try:
        body = json.dumps({"model": VLM_MODELS[attempt % len(VLM_MODELS)], "stream": False, "options": {"temperature": 0}, "messages": [{
            "role": "user", "images": [base64.b64encode(png).decode()],
            "content": "Read the captcha text in this image exactly (case-sensitive). Reply with ONLY the characters."}]}).encode()
        with urllib.request.urlopen(urllib.request.Request(f"{OLLAMA}/api/chat", body, {"Content-Type": "application/json"}), timeout=180) as r:
            return re.sub(r"\W+", "", json.load(r)["message"]["content"])
    except Exception as e:
        print("  VLM captcha error:", e); return ""


# ---- limits and pacing
BURST_WAIT = 60  # seconds to pause after a 429 "unusual request pattern" (the block was measured to clear within a minute)
MAX_DELAY = 10.0
COOLDOWN_H = 3.0  # hours an account rests after a refusal that names no time
CRASH_WAIT = 30  # seconds before resuming after an unexpected exception
ACCOUNT_BUDGET = int(os.environ.get("ACCOUNT_BUDGET", "40"))  # check-areg calls per visit
ACCOUNT_MIN_VISIT = int(os.environ.get("ACCOUNT_MIN_VISIT", "20"))  # an account with less room than this is not worth a login
# (window seconds, max calls): the site's 50/1h, 125/4h, 175/8h and 250/24h each with a margin (+2 min on the window)
CAPS = [tuple(int(x) for x in c.split(":")) for c in os.environ.get("CALL_CAPS", "3720:40,14520:112,28920:157,86520:225").split(",")]
# check-areg messages that mean "try again / log in again", NOT "this survey number has no sub-divisions"
TRANSIENT = ("session", "login", "unauthor", "not ready", "decrypt", "expired", "token", "failed to check")


def save_wb(wb, f):
    """Atomic save: a crash mid-write can never leave a truncated workbook behind."""
    tmp = f.with_name(f"{f.stem}.{os.getpid()}.tmp.xlsx"); wb.save(tmp); os.replace(tmp, f)  # pid: two processes never share a tmp


def save_json(path, data):
    tmp = path.with_name(path.name + ".tmp"); tmp.write_text(json.dumps(data)); os.replace(tmp, path)


def load_json(path):
    try: return json.loads(path.read_text())
    except (OSError, ValueError): return {}


def exit_ip(api):
    """Public IP the requests arrive from (through the proxy when one is set), or '?' if the echo failed."""
    try: return api.s.get("https://api.ipify.org", timeout=15).text.strip()[:45] or "?"
    except Exception as e: return f"? ({str(e).splitlines()[0][:60]})"


def wait_lockout():
    """Sleep until the recorded login lockout is over, without trying anything (a retry restarts the server's timer)."""
    try: until = float(LOCK_FILE.read_text())
    except (OSError, ValueError): return
    wait = until - time.time()
    if wait > 0:
        print(f"  login lockout until {time.strftime('%Y-%m-%d %H:%M', time.localtime(until))}; sleeping {wait / 60:.0f} min without trying")
        time.sleep(wait)


class SkipAccount(Exception):
    """This account cannot be used right now (login failed, OTP, ...); the pool moves on."""


class Pool:
    """This worker's accounts: which one is in use, how many calls it may still make, and when the others are free."""
    def __init__(self):
        self.pool = ACCOUNT_POOL
        self.blocked = {int(k): float(v) for k, v in load_json(COOL_FILE).items() if int(k) in self.pool}  # idx -> epoch usable again
        self.calls = {int(k): v for k, v in load_json(CALLS_FILE).items() if int(k) in self.pool}  # idx -> times of its check-areg calls
        self.pub = Api(PROXY)  # anonymous client for the public listings
        self.apis, self.api, self.idx, self.left, self.last_call, self.delay = {}, None, None, 0, 0.0, args.delay
        try: saved = int(IDX_FILE.read_text()) if IDX_FILE.exists() else self.pool[0]
        except ValueError: saved = self.pool[0]
        # resume at the saved account, or (after a lone run is split into workers) at the next one of ours past it
        self.pos = self.pool.index(saved) if saved in self.pool else next((i for i, g in enumerate(self.pool) if g > saved), 0)

    # -- call ledger: the server reports only the limit that trips, so every window is counted here
    def room(self, idx, now=None):
        now = now or time.time(); st = self.calls.get(idx, [])
        return min(cap - sum(1 for t in st if now - t < win) for win, cap in CAPS)

    def free_at(self, idx, need, now=None):
        """Epoch at which idx has room for `need` more calls in every window."""
        now = now or time.time(); st = sorted(self.calls.get(idx, [])); t_ok = now
        for win, cap in CAPS:
            live = [t for t in st if now - t < win]; k = cap - need
            if len(live) > k: t_ok = max(t_ok, live[len(live) - k - 1] + win)  # the (n-k)-th oldest call must have expired
        return t_ok

    def usable_at(self, idx, now=None): return max(self.blocked.get(idx, 0), self.free_at(idx, ACCOUNT_MIN_VISIT, now))

    def record(self, idx):
        now = time.time(); keep = max(w for w, _ in CAPS)
        self.calls[idx] = [t for t in self.calls.get(idx, []) if now - t < keep] + [now]
        save_json(CALLS_FILE, self.calls)

    def block(self, idx, seconds):
        """The server refused this account: rest it for the time it named."""
        self.blocked[idx] = time.time() + seconds; save_json(COOL_FILE, self.blocked)
        if idx == self.idx: self.left = 0

    # -- accounts
    def choose(self):
        n = len(self.pool)
        while True:
            now = time.time()
            for i in range(n):
                idx = self.pool[(self.pos + i) % n]
                if self.usable_at(idx, now) <= now: self.pos = (self.pos + i) % n; return idx
            wait = max(0, min(self.usable_at(g, now) for g in self.pool) - now) + 15
            print(f"  all {n} accounts of this worker are resting; sleeping {wait / 60:.0f} min until the first one has room")
            time.sleep(wait)

    def visit(self):
        """Move to the next account with room, get it a session, and give it its quota of calls."""
        while True:
            if self.idx is not None: self.pos = (self.pool.index(self.idx) + 1) % len(self.pool)
            self.idx = self.choose(); IDX_FILE.write_text(str(self.idx))
            quota = min(ACCOUNT_BUDGET, self.room(self.idx))
            self.api = self.apis.setdefault(self.idx, Api(PROXY, TOKEN_DIR / f"{self.idx}.json"))
            try: self.authenticate()
            except SkipAccount: continue
            self.left = quota
            print(f"Account #{self.idx}: {quota} calls this visit")
            return

    def authenticate(self):
        """Saved tokens are reused (refreshed when near expiry); log in with a captcha only when they are dead."""
        if self.api.has_tokens:
            try: self.api.maybe_refresh(); return
            except (AuthError, requests.RequestException) as e: print(f"  account #{self.idx}: tokens unusable ({str(e)[:80]}); logging in")
        self.login()

    def login(self):
        acc = ACCOUNTS[self.idx]
        while True:
            wait_lockout()
            print(f"Logging in with account #{self.idx} ({acc['mobile']}) via {'proxy' if PROXY else 'direct'}, exit IP {exit_ip(self.api)}")
            try:
                self.api.login(acc["mobile"], acc["password"], read_captcha); print("Logged in."); return
            except Locked as e:
                LOCK_FILE.write_text(str(time.time() + e.seconds)); print(f"  network lockout: {str(e)!r}")  # wait_lockout sleeps it out
            except AuthError as e:
                print(f"  account #{self.idx} skipped: {e}"); self.block(self.idx, 1800); raise SkipAccount()
            except requests.RequestException as e:
                print(f"  login network error ({str(e).splitlines()[0][:100]}), retrying in 20s"); time.sleep(20)

    def reauth(self):
        """Mid-visit trouble: refresh the tokens, else log in again; give up on the account if that fails too."""
        try: self.api.refresh(); return
        except (AuthError, requests.RequestException): pass
        try: self.login()
        except SkipAccount: self.left = 0

    def ensure(self):
        if self.api is None: self.visit()

    def before_call(self):
        while self.idx is None or self.left <= 0: self.visit()
        gap = self.delay - (time.time() - self.last_call)
        if gap > 0: time.sleep(gap)

    def after_call(self, rl):
        """A check-areg call got an answer: count it. The server's own 1 h count can say more than ours (account used elsewhere)."""
        self.last_call = time.time(); self.record(self.idx); self.left -= 1
        if rl and rl.get("window_hours") == 1 and rl.get("count", 0) >= ACCOUNT_BUDGET: self.left = 0


def api(P, path, **params): return P.pub.public(path, **params)


def geom_center(data):
    """(lat, lon) of the bounding-box centre of a GeoJSON FeatureCollection, or None."""
    pts = []
    def walk(c):
        if c and isinstance(c[0], (int, float)): pts.append(c[:2])
        else: [walk(x) for x in c or []]
    for f in (data or {}).get("features", []): walk((f.get("geometry") or {}).get("coordinates"))
    if not pts: return None
    lon = (min(p[0] for p in pts) + max(p[0] for p in pts)) / 2; lat = (min(p[1] for p in pts) + max(p[1] for p in pts)) / 2
    if abs(lon) > 180 or abs(lat) > 90:  # defensive: data arrived in Web Mercator (EPSG:3857)
        import math
        lon, lat = lon / 20037508.34 * 180, math.degrees(2 * math.atan(math.exp(lat / 6378137)) - math.pi / 2)
    return round(lat, 6), round(lon, 6)


def village_center(P, d, t, v):
    """(lat, lon) of the village centre as the site's map shows it, or None if it could not be fetched."""
    for attempt in range(3):
        P.ensure()
        try:
            status, r = P.api.secure("land/get-geom", {"case": "revenue_village", "code_type": "revenue", "district_code": d, "taluk_code": t,
                                                       "village_code": v, "survey_number": None, "sub_division_number": None})
        except (AuthError, requests.RequestException) as e:
            print(f"  get-geom error ({str(e).splitlines()[0][:100]}); retrying"); time.sleep(10); P.reauth(); continue
        if r.get("success") == 1:
            c = geom_center(r.get("data"))
            if c: return c
            print(f"  get-geom: no geometry for village {v}"); return None
        msg = str(r.get("message", "")).lower()
        if "slow down" in msg or "unusual" in msg:
            print(f"  burst limiter tripped on get-geom; waiting {BURST_WAIT}s"); time.sleep(BURST_WAIT); continue
        if r.get("rate_limited") or "limit" in msg:
            wait = (r.get("rate_limit_status") or {}).get("retry_after_seconds") or parse_wait(msg) or COOLDOWN_H * 3600
            print(f"  get-geom rate limited ({r.get('message')}); resting account #{P.idx}"); P.block(P.idx, wait + 60); P.visit(); continue
        print(f"  get-geom failed for village {v}: {r.get('message')}"); time.sleep(3)
    return None


COORDS_RE = r"_-?\d+\.\d+_-?\d+\.\d+\.xlsx$"  # "_<lat>_<long>.xlsx"


def out_name(dn, tn, vn, center):
    stem = f"Village_map_of_{safe(dn)}_{safe(tn)}_{safe(vn)}"
    return OUT / (f"{stem}_{center[0]}_{center[1]}.xlsx" if center else f"{stem}.xlsx")


def has_coords(f): return re.search(COORDS_RE, f.name) is not None


def find_file(stem):
    """The village's workbook, with or without coordinates in the name, or None. A regex (not a glob) so that
    'Village_map_of_X_Y_Peria' can never match 'Village_map_of_X_Y_Peria_Obulapuram_<lat>_<long>.xlsx'."""
    pat = re.compile("^" + re.escape(stem) + "(" + COORDS_RE + "|\\.xlsx$)")
    return next((f for f in sorted(OUT.glob(f"{stem}*.xlsx")) if pat.match(f.name)), None)


SEEN_MSGS = set()


def subdivs(P, d, t, v, s):
    """Returns (list of sub-division numbers, rate_limit_status or {}). Waits for / switches accounts as needed."""
    fails = 0
    while True:
        P.before_call()
        try:
            status, r = P.api.secure("land/check-areg", {"district_code": d, "taluk_code": t, "village_code": v, "survey_number": s,
                                                         "sub_division_number": "jjj", "area_type": "rural"})
        except (AuthError, requests.RequestException) as e:
            fails += 1; print(f"  check-areg error ({str(e).splitlines()[0][:120]}), retrying in 10s"); time.sleep(10)
            if fails % 3 == 0: P.reauth()
            continue
        rl = r.get("rate_limit_status") or {}
        msg = str(r.get("message", "") or r.get("error", "")).lower()
        if status == 200 and r.get("success") == 2:
            P.after_call(rl); return [x["subdiv_no"] for x in r.get("data", [])], rl
        if status == 429 or r.get("rate_limited") or "limit" in msg or "too many" in msg:
            P.after_call(rl)
            if "slow down" in msg or "unusual" in msg:  # burst detector: wait it out, then go slower
                P.delay = min(P.delay + 1, MAX_DELAY)
                print(f"  burst limiter tripped ({msg[:60]}); waiting {BURST_WAIT}s, delay now {P.delay}s")
                time.sleep(BURST_WAIT); continue
            wait = rl.get("retry_after_seconds") or parse_wait(msg) or COOLDOWN_H * 3600
            print(f"  limit reached ({r.get('message')}); resting account #{P.idx} {wait / 3600:.1f} h"); P.block(P.idx, wait + 60); continue
        if (status is not None and status != 200) or any(k in msg for k in TRANSIENT):  # 5xx, dead session...
            fails += 1; print(f"  check-areg transient failure (HTTP {status}: {msg[:80]}), retrying"); time.sleep(10)
            if fails % 3 == 0: P.reauth()
            continue
        P.after_call(rl)
        key = (r.get("success"), msg)
        if key not in SEEN_MSGS:  # log each distinct non-data response once so empty results are not silent
            SEEN_MSGS.add(key); print(f"  note: check-areg returned no sub-divisions for survey {s}: {str(r)[:200]}")
        return [], rl  # no sub-divisions for this survey number


def pending_rows(ws):
    """[(survey_number, row)] whose sub-divisions cell is still blank = not fetched yet."""
    return [(str(ws.cell(r, 1).value), r) for r in range(2, ws.max_row + 1) if ws.cell(r, 2).value in (None, "")]


def village(P, dc, dn, tc, tn, vc, vn):
    """Scrape (or finish) one village. Returns False if another process is writing its workbook right now."""
    stem = f"Village_map_of_{safe(dn)}_{safe(tn)}_{safe(vn)}"
    for junk in OUT.glob(f"{stem}*.tmp.xlsx"):
        if time.time() - junk.stat().st_mtime > ACTIVE_WINDOW: junk.unlink()  # left over from a crash mid-save
    f = find_file(stem)
    if f:
        if time.time() - f.stat().st_mtime < ACTIVE_WINDOW:
            wb = load_workbook(f, read_only=True)
            busy = bool(pending_rows(wb.active)); wb.close()
            if busy:
                print(f"{dn} / {tn} / {vn}: workbook saved {int(time.time() - f.stat().st_mtime)}s ago by another process; will revisit"); return False
        wb = load_workbook(f); ws = wb.active
        if not pending_rows(ws) and has_coords(f): return True  # finished
    if f is None or not has_coords(f):  # need the centre (new village, or old file without coordinates)
        center = village_center(P, dc, tc, vc); time.sleep(P.delay)
        if f is None:
            surveys = [x["survey_number"] for x in api(P, "admin_master_survey_number", district_code=dc, taluk_code=tc,
                       revenue_village_code=vc, area_type="rural", data_type="cadastral", request_type="survey_number")]
            f = out_name(dn, tn, vn, center)
            wb = Workbook(); ws = wb.active; ws.title = "data"
            ws.append(["surveynumber", "sub-divisions"])
            for s in surveys: ws.append([s, None])
            save_wb(wb, f)
        elif center:
            new = out_name(dn, tn, vn, center); f.rename(new); f = new
            print(f"{dn} / {tn} / {vn}: added centre {center} to file name")
    todo = pending_rows(ws)
    total = ws.max_row - 1
    print(f"{dn} / {tn} / {vn}: {total} survey numbers ({total - len(todo)} done) -> {f.name}")
    for i, (s, row) in enumerate(todo):
        subs, rl = subdivs(P, dc, tc, vc, s)
        ws.cell(row, 2).value = ", ".join(subs) or "-"   # never leave a fetched row blank
        save_wb(wb, f)
        if i % 25 == 0: print(f"  {total - len(todo) + i}/{total} account #{P.idx} ({P.left} calls left this visit) rate_limit={rl}")
    print("  finished", f.name)
    return True


def run(P):
    print(f"worker {WK}/{NWK}: accounts {ACCOUNT_POOL}")
    deferred = []
    for dc, dn in [(x["district_code"], x["district_english_name"]) for x in api(P, "admin_master_district", request_type="district")
                   if match(x["district_english_name"], args.district)]:
        for tx in api(P, "admin_master_taluk", district_code=dc, request_type="taluk"):
            tc, tn = tx["taluk_code"], tx["taluk_english_name"]
            if not match(tn, args.taluk): continue
            for vi, vx in enumerate(api(P, "admin_master_village", district_code=dc, taluk_code=tc, request_type="revenue_village")):
                if vi % NWK != WK: continue  # another worker's village
                vc, vn = vx["village_code"], vx["village_english_name"]
                if not match(vn, args.village): continue
                if not village(P, dc, dn, tc, tn, vc, vn): deferred.append((dc, dn, tc, tn, vc, vn))
            if deferred:  # villages another process was writing: try them again after every taluk
                deferred = [v for v in deferred if not village(P, *v)]
    while deferred:
        print(f"{len(deferred)} village(s) were being written by another process; re-checking in 5 min")
        time.sleep(300)
        deferred = [v for v in deferred if not village(P, *v)]


def main():
    """Crash-recovery wrapper: whatever goes wrong, wait a bit and resume from the saved progress."""
    P = Pool()
    print(f"proxy method {METHOD}: {'direct' if not PROXY else re.sub(r'//[^@/]*@', '//***@', PROXY)}; delay {P.delay}s; caps {CAPS}")
    while True:
        try:
            run(P); print("All done."); return
        except KeyboardInterrupt:
            print("interrupted; progress is saved in the workbooks, just start again to resume"); return
        except Exception:
            traceback.print_exc()
            print(f"  unexpected error; resuming in {CRASH_WAIT}s"); time.sleep(CRASH_WAIT)


if __name__ == "__main__":
    main()
